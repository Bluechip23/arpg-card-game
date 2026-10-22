#!/usr/bin/env python3
"""Rebuild docs/SLOT_SHARING.csv, .xlsx and .png from the game's cards.

For every pair of slot labels: how many cards carry both (a card with both
labels can go in either kind of item). The cards and their labels come from
the game itself (tools/dump_slot_labels.gd -> Card.get_all_cards), so the
three files always agree with the code. Run after slot labels change:
  python3 tools/build_slot_sharing.py
Needs openpyxl and matplotlib (tool dependencies, not the game's).
"""
import json
import os
import subprocess
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ORDER = ["Sword", "Axe", "Dagger", "Hammer", "Spear", "Wand", "Staff", "Arrow",
         "Buckler", "Bulwark", "Crown", "Fist", "Swift", "Pocket", "Gem"]
OUT = "docs/SLOT_SHARING"


def load_cards():
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "cards.json")
        subprocess.run(["godot", "--headless", "--path", ".", "--script",
                        "tools/dump_slot_labels.gd", "--", path],
                       check=True, capture_output=True)
        with open(path) as f:
            return json.load(f)


def matrix(cards):
    sets = [set(c["labels"]) for c in cards if c["labels"]]
    m = {a: {b: sum(1 for s in sets if a in s and b in s) for b in ORDER} for a in ORDER}
    total = {a: sum(1 for s in sets if a in s) for a in ORDER}
    only = {a: sum(1 for s in sets if s == {a}) for a in ORDER}
    return m, total, only, len(sets)


def write_csv(m, total, only):
    lines = ["Slot," + ",".join(ORDER) + ",Cards with this slot,Only this slot"]
    for a in ORDER:
        row = [a] + ["" if a == b else str(m[a][b]) for b in ORDER] + [str(total[a]), str(only[a])]
        lines.append(",".join(row))
    with open(OUT + ".csv", "w") as f:
        f.write("\n".join(lines) + "\n")


def green(n, top):
    # White to the chart's green, by share of the busiest pair.
    t = 0 if top == 0 else min(1.0, n / top)
    r, g, b = 255 - int(t * (255 - 99)), 255 - int(t * (255 - 190)), 255 - int(t * (255 - 123))
    return r, g, b


def write_xlsx(cards, m, total, only, n_slottable):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    ws["A1"] = "Number of cards that can be slotted into both slots"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Each cell counts the cards whose slot labels include both the row's and the "
                "column's slot. Source: every card the game builds (Card.create_by_id), "
                "%d cards, %d of them slottable. Tome was folded into Wand (2026-10-20)."
                % (len(cards), n_slottable))
    head = ["Slot"] + ORDER + ["Cards with this slot", "Only this slot"]
    for j, h in enumerate(head):
        c = ws.cell(row=4, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    top = max(m[a][b] for a in ORDER for b in ORDER if a != b)
    for i, a in enumerate(ORDER):
        ws.cell(row=5 + i, column=1, value=a).font = Font(bold=True)
        for j, b in enumerate(ORDER):
            c = ws.cell(row=5 + i, column=2 + j)
            c.alignment = Alignment(horizontal="center")
            if a == b:
                c.fill = PatternFill("solid", fgColor="808080")
                continue
            c.value = m[a][b]
            r, g, bl = green(m[a][b], top)
            c.fill = PatternFill("solid", fgColor="%02X%02X%02X" % (r, g, bl))
        ws.cell(row=5 + i, column=2 + len(ORDER), value=total[a]).font = Font(bold=True)
        ws.cell(row=5 + i, column=3 + len(ORDER), value=only[a])
    ws.column_dimensions["A"].width = 10
    cs = wb.create_sheet("Cards")
    cs.append(["Card", "Card id", "Slots", "Slot count", "Engrave only"] + ORDER)
    for c in sorted((c for c in cards if c["labels"]), key=lambda c: (-len(c["labels"]), c["name"])):
        cs.append([c["name"], c["id"], " / ".join(c["labels"]), len(c["labels"]),
                   "yes" if c["engrave"] else None] + ["✓" if s in c["labels"] else None for s in ORDER])
    for col in ("A", "B", "C"):
        cs.column_dimensions[col].width = 28
    wb.save(OUT + ".xlsx")


def write_png(cards, m, total, n_slottable):
    head = ["Slot"] + ORDER + ["Total"]
    rows = []
    for a in ORDER:
        rows.append([a] + ["" if a == b else str(m[a][b]) for b in ORDER] + [str(total[a])])
    top = max(m[a][b] for a in ORDER for b in ORDER if a != b)
    fig, ax = plt.subplots(figsize=(10.5, 6.2), dpi=100)
    ax.axis("off")
    t = ax.table(cellText=rows, colLabels=head, loc="center", cellLoc="center")
    t.auto_set_font_size(False)
    t.set_fontsize(10)
    t.auto_set_column_width(col=list(range(len(head))))
    t.scale(1, 1.55)
    for (i, j), cell in t.get_celld().items():
        cell.set_edgecolor("#bbbbbb")
        if i == 0 or j == 0:
            cell.set_facecolor("#e3e3e3")
            continue
        if j == len(head) - 1:
            cell.set_facecolor("#f2f2f2")
            cell.get_text().set_fontweight("bold")
            continue
        a, b = ORDER[i - 1], ORDER[j - 1]
        if a == b:
            cell.set_facecolor("#808080")
            continue
        n = m[a][b]
        r, g, bl = green(n, top)
        cell.set_facecolor((r / 255, g / 255, bl / 255))
        if n == 0:
            cell.get_text().set_color("#aaaaaa")
    ax.set_title("Number of cards that can be slotted into both slots", loc="left",
                 color="#2e6b1f", fontweight="bold", fontsize=14)
    fig.text(0.02, 0.1, "Total = cards that can go in that slot. Counted from every card the game "
             "builds (%d cards, %d slottable). Tome was folded into Wand." % (len(cards), n_slottable),
             fontsize=8, color="#555555")
    fig.savefig(OUT + ".png", bbox_inches="tight")


def main():
    cards = load_cards()
    m, total, only, n = matrix(cards)
    write_csv(m, total, only)
    write_xlsx(cards, m, total, only, n)
    write_png(cards, m, total, n)
    print("slot sharing: %d cards, %d slottable -> %s.csv/.xlsx/.png" % (len(cards), n, OUT))


if __name__ == "__main__":
    main()
